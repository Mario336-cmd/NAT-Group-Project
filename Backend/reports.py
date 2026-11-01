# Build local PDF and Word reports from a saved browser result.

from datetime import datetime, timezone, timedelta
from functools import lru_cache
from io import BytesIO
import os
from pathlib import Path
import re
from threading import Lock
from typing import Literal
from urllib.parse import urlsplit
from xml.sax.saxutils import escape

from pydantic import BaseModel, ConfigDict, Field, FiniteFloat, field_validator, model_validator


ROOT = Path(__file__).resolve().parents[1]
CONFIDENCE_NOTE = "Confidence is a model score, not a guarantee of accuracy."
FONT_LOCK = Lock()


class ReportModel(BaseModel):
    model_config = ConfigDict(extra="ignore", strict=True, allow_inf_nan=False)


class Prediction(ReportModel):
    label: Literal["SUPPORTS", "REFUTES"]
    confidence: float = Field(ge=0, le=1)
    chunks: int | None = Field(default=None, ge=1)
    probabilities: dict[Literal["SUPPORTS", "REFUTES"], FiniteFloat] = Field(default_factory=dict)


class Explanation(ReportModel):
    status: Literal["ready", "unavailable", "error"]
    summary: str = Field(default="", max_length=5000)
    citations: list[str] = Field(default_factory=list, max_length=10)
    reason: str | None = Field(default=None, max_length=2000)


class Screening(ReportModel):
    status: Literal["accepted", "invalid_input", "out_of_scope", "unavailable"]
    reason: str | None = Field(default=None, max_length=2000)


class Passage(ReportModel):
    id: str = Field(pattern=r"^E[1-9]\d*$", max_length=8)
    page: str = Field(max_length=2000)
    text: str = Field(max_length=200000)
    url: str = Field(max_length=4096)
    pid: int | None = None
    sentence_id: int | None = None
    score: FiniteFloat | None = None
    emb_score: FiniteFloat | None = None


class Result(ReportModel):
    claim: str = Field(min_length=1, max_length=4000)
    status: Literal["prediction", "no_evidence", "rejected", "screening_unavailable"]
    top_k: int = Field(ge=1, le=10)
    prediction: Prediction | None
    explanation: Explanation
    screening: Screening | None = None
    evidence: list[Passage] = Field(max_length=10)
    notice: str = Field(default="", max_length=5000)
    timings: dict[str, FiniteFloat] = Field(default_factory=dict, max_length=20)

    @model_validator(mode="after")
    def consistent_result(self):
        ids = [item.id for item in self.evidence]
        if (len(set(ids)) != len(ids) or len(ids) > self.top_k
                or not set(self.explanation.citations) <= set(ids)
                or any(len(name) > 80 for name in self.timings)):
            raise ValueError("Inconsistent report evidence")
        if self.status == "prediction" and (self.prediction is None or not ids):
            raise ValueError("Incomplete report prediction")
        if self.status == "no_evidence" and (self.prediction is not None or ids):
            raise ValueError("Inconsistent no-evidence result")
        if self.status == "prediction" and self.screening is not None:
            if self.screening.status != "accepted" or self.explanation.status != "ready":
                raise ValueError("Prediction was not accepted by screening")
        if self.status in ("rejected", "screening_unavailable"):
            allowed = ("invalid_input", "out_of_scope") if self.status == "rejected" else ("unavailable",)
            if (self.prediction is not None or self.screening is None or self.screening.status not in allowed
                    or not (self.screening.reason or "").strip() or self.explanation.status == "ready"
                    or self.explanation.summary or self.explanation.citations):
                raise ValueError("Inconsistent screening result")
        return self


class ReportRecord(ReportModel):
    id: str = Field(pattern=r"^[A-Za-z0-9_-]{1,128}$")
    saved_at: str = Field(min_length=20, max_length=40)
    result: Result

    @field_validator("saved_at")
    @classmethod
    def valid_completion_time(cls, value):
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            raise ValueError("A completion timezone is required")
        return value


def clean_text(value):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "\ufffd", str(value))


def safe_source_url(value):
    try:
        parsed = urlsplit(value)
        if (parsed.scheme in ("https", "http") and parsed.hostname
                and (parsed.hostname == "wikipedia.org" or parsed.hostname.endswith(".wikipedia.org"))
                and not parsed.username and not parsed.password):
            return value
    except ValueError:
        pass
    return None


def report_blocks(record):
    result = record.result
    completed = datetime.fromisoformat(record.saved_at.replace("Z", "+00:00"))
    completed = completed.astimezone(timezone(timedelta(hours=2))).strftime("%d %B %Y, %H:%M CAT")
    source_numbers = {item.id: index for index, item in enumerate(result.evidence, 1)}

    def citation_label(source_id):
        return "[" + str(source_numbers.get(source_id, source_id)) + "]"

    def readable_summary(text):
        def replace_group(match):
            ids = re.split(r",\s*", match.group(1))
            if not all(source_id in source_numbers for source_id in ids):
                return match.group(0)
            return "[" + ", ".join(str(source_numbers[source_id]) for source_id in ids) + "]"
        return re.sub(r"\[(E[1-9]\d*(?:,\s*E[1-9]\d*)*)\]", replace_group, text)

    blocks = [("title", "Claim verification report"), ("meta", "Completed " + completed),
              ("heading", "Claim"), ("body", result.claim)]
    blocked = result.status in ("rejected", "screening_unavailable")
    if blocked:
        label = "Screening unavailable" if result.status == "screening_unavailable" else (
            "Check your input" if result.screening.status == "invalid_input" else "Enter a factual claim"
        )
        blocks.extend([("heading", "Input screening"), ("body", label), ("body", result.screening.reason)])
    elif result.prediction is not None:
        label = "Supports" if result.prediction.label == "SUPPORTS" else "Refutes"
        blocks.extend([("heading", "Model result"), ("body", label),
                       ("body", f"Model confidence: {result.prediction.confidence:.1%}")])
        if result.screening is None:
            blocks.append(("meta", "Saved before input screening was added."))
    else:
        blocks.extend([("heading", "Model result"), ("body", "No passages were found for this claim.")])
    if not blocked and result.status != "no_evidence":
        blocks.append(("heading", "What the evidence says"))
        if result.explanation.status == "ready":
            summary = readable_summary(result.explanation.summary) or "No explanation was returned."
            blocks.extend(("body", paragraph.strip()) for paragraph in re.split(r"\n\s*\n", summary) if paragraph.strip())
            if result.explanation.citations:
                blocks.append(("meta", "Sources: " + ", ".join(citation_label(item) for item in result.explanation.citations)))
        else:
            blocks.append(("body", "Explanation unavailable. " + (result.explanation.reason or "Please try again later.")))
    if result.evidence:
        count = len(result.evidence)
        heading = "Passages reviewed" if blocked else "Evidence used"
        heading += " · " + str(count) + (" passage" if count == 1 else " passages")
        blocks.append(("heading", heading))
        for item in result.evidence:
            blocks.extend([("source", citation_label(item.id) + " " + item.page.replace("_", " ")),
                           ("body", item.text), ("url", item.url)])
    if result.prediction is not None:
        blocks.append(("meta", CONFIDENCE_NOTE))
    return [(kind, clean_text(text)) for kind, text in blocks]


def filename(record, extension):
    stamp = datetime.fromisoformat(record.saved_at.replace("Z", "+00:00")).astimezone(timezone.utc).strftime("%Y-%m-%dT%H-%M-%SZ")
    return "truecheck-" + stamp + "-" + record.id[:64] + "." + extension


@lru_cache(maxsize=1)
def pdf_fonts():
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    candidates = [
        (Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/arial.ttf", "arialbd.ttf"),
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), "DejaVuSans-Bold.ttf"),
        (Path(reportlab.__file__).parent / "fonts/Vera.ttf", "VeraBd.ttf"),
    ]
    with FONT_LOCK:
        regular, bold = next((path, path.with_name(bold)) for path, bold in candidates if path.is_file() and path.with_name(bold).is_file())
        pdfmetrics.registerFont(TTFont("TrueCheckBody", str(regular)))
        pdfmetrics.registerFont(TTFont("TrueCheckBold", str(bold)))
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "TrueCheckBody", "TrueCheckBold"


def pdf_report(record):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    normal, bold = pdf_fonts()
    stream = BytesIO()
    document = SimpleDocTemplate(stream, pagesize=letter, rightMargin=54, leftMargin=54,
                                 topMargin=66, bottomMargin=54, title="TrueCheck claim verification report", author="TrueCheck")
    styles = {
        "title": ParagraphStyle("title", fontName=bold, fontSize=22, leading=28, textColor=colors.black, spaceAfter=14, keepWithNext=True),
        "heading": ParagraphStyle("heading", fontName=bold, fontSize=13, leading=18, textColor=colors.black, spaceBefore=15, spaceAfter=7, keepWithNext=True),
        "source": ParagraphStyle("source", fontName=bold, fontSize=11, leading=16, spaceBefore=9, spaceAfter=5, keepWithNext=True),
        "body": ParagraphStyle("body", fontName=normal, fontSize=12, leading=18, spaceAfter=8, alignment=TA_LEFT, splitLongWords=True),
        "meta": ParagraphStyle("meta", fontName=normal, fontSize=10, leading=15, textColor=colors.HexColor("#596579"), spaceAfter=5, splitLongWords=True),
        "url": ParagraphStyle("url", fontName=normal, fontSize=10, leading=15, textColor=colors.HexColor("#1d4ed8"), spaceAfter=7, splitLongWords=True),
    }

    def markup(text):
        value = escape(text).replace("\n", "<br/>")
        # CJK uses the PDF standard's Unicode font while Latin text stays embedded.
        return re.sub(r"[\u2e80-\ua000\uf900-\ufaff]+", lambda match: '<font name="STSong-Light">' + match.group() + '</font>', value)

    flow = []
    for kind, text in report_blocks(record):
        value = markup(text)
        if kind == "url" and safe_source_url(text):
            value = '<link href="' + escape(text, {'"': '&quot;'}) + '">' + value + '</link>'
        flow.append(Paragraph(value, styles[kind]))

    def furniture(canvas, doc):
        canvas.saveState()
        logo = ROOT / "Frontend/assets/logo.png"
        if logo.is_file():
            canvas.drawImage(str(logo), 54, 745, width=19, height=19, mask="auto")
        canvas.setFont(bold, 11)
        canvas.drawString(81, 750, "TrueCheck")
        canvas.setFont(normal, 10)
        canvas.setFillColor(colors.HexColor("#596579"))
        caption = "Prediction with retrieved evidence" if record.result.prediction is not None else (
            "Input screening result" if record.result.status in ("rejected", "screening_unavailable") else "No evidence result"
        )
        canvas.drawString(54, 32, caption)
        canvas.drawRightString(558, 32, f"Page {doc.page}")
        canvas.restoreState()

    document.build(flow, onFirstPage=furniture, onLaterPages=furniture)
    return stream.getvalue()


def word_report(record):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.85), Inches(.75)
    section.left_margin = section.right_margin = Inches(.75)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(8)
    for name, size in (("Title", 22), ("Heading 1", 14), ("Heading 2", 11)):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
        borders = style.element.find("w:pPr/w:pBdr", namespaces=style.element.nsmap)
        if borders is not None:
            borders.getparent().remove(borders)
    for style in (normal, document.styles["Title"], document.styles["Heading 1"], document.styles["Heading 2"]):
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + attribute), None)
    header = section.header.paragraphs[0]
    logo = ROOT / "Frontend/assets/logo.png"
    if logo.is_file():
        header.add_run().add_picture(str(logo), width=Inches(.23))
    header.add_run("  TrueCheck").bold = True
    for kind, text in report_blocks(record):
        style = {"title": "Title", "heading": "Heading 1", "source": "Heading 2"}.get(kind)
        paragraph = document.add_paragraph(style=style)
        if kind == "url" and safe_source_url(text):
            from docx.opc.constants import RELATIONSHIP_TYPE
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), paragraph.part.relate_to(text, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "1D4ED8")
            properties.append(color)
            run.append(properties)
            content = OxmlElement("w:t")
            content.text = text
            run.append(content)
            link.append(run)
            paragraph._p.append(link)
        else:
            paragraph.add_run(text)
        if kind in ("meta", "url"):
            paragraph.paragraph_format.space_after = Pt(5)
            for run in paragraph.runs:
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor.from_string("596579" if kind == "meta" else "1D4ED8")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("TrueCheck | Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.author = "TrueCheck"
    document.core_properties.title = "TrueCheck claim verification report"
    document.core_properties.comments = "Generated locally from the selected verification result."
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
